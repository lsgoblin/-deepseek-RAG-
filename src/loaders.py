"""Document loading utilities."""

from __future__ import annotations

import hashlib
import logging
from pathlib import Path
from typing import Iterable

try:
    from src.schemas import SourceDocument
except ModuleNotFoundError:  # pragma: no cover - script execution fallback
    from schemas import SourceDocument


LOGGER = logging.getLogger(__name__)


SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".docx", ".xlsx")
DEFAULT_EXCLUDED_DIR_NAMES = {"raw"}


def discover_documents(
    data_dir: Path,
    max_docs: int | None = None,
    extensions: Iterable[str] = SUPPORTED_EXTENSIONS,
    excluded_dir_names: Iterable[str] = DEFAULT_EXCLUDED_DIR_NAMES,
) -> list[Path]:
    """Return candidate document paths from the data directory."""

    normalized_extensions = {ext.lower() for ext in extensions}
    excluded_parts = {part.lower() for part in excluded_dir_names}
    discovered = sorted(
        file_path
        for file_path in data_dir.rglob("*")
        if file_path.is_file()
        and file_path.suffix.lower() in normalized_extensions
        and not any(part.lower() in excluded_parts for part in file_path.relative_to(data_dir).parts[:-1])
    )

    if max_docs is not None:
        return discovered[:max_docs]
    return discovered


def _build_doc_id(file_path: Path) -> str:
    """Create a stable document identifier from the file path."""

    digest = hashlib.md5(str(file_path.resolve()).encode("utf-8")).hexdigest()[:12]
    return f"{file_path.stem}-{digest}"


def _read_text_file(file_path: Path) -> str:
    """Read plain-text content with a few common encoding fallbacks."""

    candidate_encodings = ("utf-8", "utf-8-sig", "gb18030")
    last_error: UnicodeDecodeError | None = None

    for encoding in candidate_encodings:
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError as exc:
            last_error = exc

    if last_error is not None:
        raise last_error
    return ""


def _read_pdf_file(file_path: Path) -> str:
    """Extract text from a PDF file."""

    try:
        from pypdf import PdfReader
    except ModuleNotFoundError as exc:  # pragma: no cover - depends on runtime env
        raise RuntimeError("Missing dependency `pypdf`. Install requirements first.") from exc

    reader = PdfReader(str(file_path))
    return "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _read_docx_file(file_path: Path) -> str:
    """Extract text from a DOCX file."""

    try:
        from docx import Document
    except ModuleNotFoundError as exc:  # pragma: no cover - depends on runtime env
        raise RuntimeError(
            "Missing dependency `python-docx`. Install requirements first."
        ) from exc

    document = Document(str(file_path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()


def _normalize_spreadsheet_cell(value: object) -> str:
    """Normalize a spreadsheet cell into clean text."""

    if value is None:
        return ""

    text = str(value).replace("\ufeff", "").strip()
    if text.lower() == "nan":
        return ""
    return text


def _detect_header_row(rows: list[list[str]]) -> int:
    """Pick the most likely header row from the top of a worksheet."""

    best_index = 0
    best_score = 0
    for index, row in enumerate(rows[:10]):
        non_empty_cells = sum(1 for cell in row if cell)
        if non_empty_cells >= 2 and non_empty_cells > best_score:
            best_index = index
            best_score = non_empty_cells

    return best_index


def _deduplicate_headers(headers: list[str]) -> list[str]:
    """Ensure worksheet headers are non-empty and unique."""

    seen: dict[str, int] = {}
    normalized_headers: list[str] = []

    for index, raw_header in enumerate(headers, start=1):
        base_name = raw_header or f"column_{index}"
        occurrence = seen.get(base_name, 0)
        seen[base_name] = occurrence + 1
        if occurrence == 0:
            normalized_headers.append(base_name)
        else:
            normalized_headers.append(f"{base_name}_{occurrence + 1}")

    return normalized_headers


def _format_sheet_row(
    *,
    sheet_name: str,
    sheet_title: str,
    headers: list[str],
    values: list[str],
    row_number: int,
) -> str:
    """Render one worksheet row into retrieval-friendly text."""

    row_data = {
        header: value
        for header, value in zip(headers, values, strict=False)
        if header and value
    }
    if not row_data:
        return ""

    prompt_parts = [
        value
        for header, value in row_data.items()
        if header not in {"序号", "id", "ID", "编号"}
    ]

    lines = [
        f"工作表: {sheet_name}",
        f"表内行号: {row_number}",
    ]
    if sheet_title:
        lines.append(f"表格主题: {sheet_title}")

    for header, value in row_data.items():
        lines.append(f"{header}: {value}")

    if prompt_parts:
        lines.append(f"组合提示词: {'，'.join(prompt_parts)}")

    return "\n".join(lines)


def _extract_xlsx_rows(file_path: Path) -> list[dict[str, object]]:
    """Extract normalized row records from an XLSX file."""

    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:  # pragma: no cover - depends on runtime env
        raise RuntimeError("Missing dependency `openpyxl`. Install requirements first.") from exc

    workbook = load_workbook(str(file_path), read_only=True, data_only=True)
    extracted_rows: list[dict[str, object]] = []

    for worksheet in workbook.worksheets:
        rows = [
            [_normalize_spreadsheet_cell(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue

        header_row_index = _detect_header_row(rows)
        sheet_title = " | ".join(cell for row in rows[:header_row_index] for cell in row if cell)
        headers = _deduplicate_headers(rows[header_row_index])
        body_rows = rows[header_row_index + 1 :]

        for row_offset, row in enumerate(body_rows, start=header_row_index + 2):
            padded_row = row + [""] * max(0, len(headers) - len(row))
            rendered = _format_sheet_row(
                sheet_name=worksheet.title,
                sheet_title=sheet_title,
                headers=headers,
                values=padded_row,
                row_number=row_offset,
            )
            if rendered:
                row_data = {
                    header: value
                    for header, value in zip(headers, padded_row, strict=False)
                    if header and value
                }
                extracted_rows.append(
                    {
                        "sheet_name": worksheet.title,
                        "sheet_title": sheet_title,
                        "row_number": row_offset,
                        "row_data": row_data,
                        "content": rendered,
                    }
                )

    return extracted_rows


def _read_xlsx_file(file_path: Path) -> str:
    """Extract worksheet rows from an XLSX file as normalized text."""

    extracted_rows = _extract_xlsx_rows(file_path)
    return "\n\n".join(str(item["content"]) for item in extracted_rows).strip()


def _load_xlsx_documents(file_path: Path) -> list[SourceDocument]:
    """Load one XLSX file as multiple row-level knowledge documents."""

    extracted_rows = _extract_xlsx_rows(file_path)
    documents: list[SourceDocument] = []

    for item in extracted_rows:
        row_number = int(item["row_number"])
        sheet_name = str(item["sheet_name"])
        row_data = dict(item["row_data"])
        documents.append(
            SourceDocument(
                doc_id=f"{_build_doc_id(file_path)}-{sheet_name}-{row_number:04d}",
                source_path=file_path,
                source_name=f"{file_path.name}::{sheet_name}::row-{row_number}",
                file_type="xlsx",
                content=str(item["content"]).strip(),
                metadata={
                    "source": file_path.name,
                    "sheet_name": sheet_name,
                    "row_number": row_number,
                    "row_data": row_data,
                },
            )
        )

    return documents


def load_document(file_path: Path) -> SourceDocument:
    """Load one document and normalize it as `SourceDocument`."""

    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {file_path.suffix}")

    if suffix in {".txt", ".md"}:
        content = _read_text_file(file_path)
    elif suffix == ".pdf":
        content = _read_pdf_file(file_path)
    elif suffix == ".docx":
        content = _read_docx_file(file_path)
    else:
        content = _read_xlsx_file(file_path)

    return SourceDocument(
        doc_id=_build_doc_id(file_path),
        source_path=file_path,
        source_name=file_path.name,
        file_type=suffix.lstrip("."),
        content=content.strip(),
        metadata={"source": file_path.name},
    )


def load_documents(data_dir: Path, max_docs: int | None = None) -> list[SourceDocument]:
    """Load all supported documents from a knowledge-base directory."""

    loaded_documents: list[SourceDocument] = []
    discovered_paths = discover_documents(data_dir=data_dir, max_docs=None)
    for file_path in discovered_paths:
        try:
            if file_path.suffix.lower() == ".xlsx":
                loaded_documents.extend(_load_xlsx_documents(file_path))
            else:
                loaded_documents.append(load_document(file_path))
        except Exception as exc:  # pragma: no cover - exercised by runtime failures
            LOGGER.warning("Skipping document %s: %s", file_path, exc)

        if max_docs is not None and len(loaded_documents) >= max_docs:
            return loaded_documents[:max_docs]

    return loaded_documents
